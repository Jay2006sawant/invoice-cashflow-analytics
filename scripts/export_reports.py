"""
Export Tableau-ready CSVs and a stakeholder Excel summary report.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DB_PATH = PROJECT_ROOT / "database" / "invoices.db"
EXPORT_DIR = PROJECT_ROOT / "exports"
SQL_DIR = PROJECT_ROOT / "sql"
AS_OF_DATE = "2026-05-10"


def run_query(sql_path: Path) -> pd.DataFrame:
    sql = sql_path.read_text(encoding="utf-8")
    with sqlite3.connect(DB_PATH) as conn:
        return pd.read_sql_query(sql, conn)


def export_tableau_csvs() -> None:
    exports = {
        "dso_by_month.csv": SQL_DIR / "dso_by_month.sql",
        "aging_buckets.csv": SQL_DIR / "aging_buckets.sql",
        "top_late_payers.csv": SQL_DIR / "late_payment_rate_by_customer.sql",
        "cash_inflow_trends.csv": SQL_DIR / "cash_inflow_trends.sql",
        "dso_by_segment.csv": SQL_DIR / "dso_by_segment.sql",
    }

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    for filename, sql_path in exports.items():
        df = run_query(sql_path)
        output = EXPORT_DIR / filename
        df.to_csv(output, index=False)
        print(f"Exported {filename} ({len(df)} rows)")


def style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def write_sheet(wb: Workbook, title: str, df: pd.DataFrame, start_row: int = 1) -> None:
    ws = wb.create_sheet(title=title)
    for row_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        ws.append(row)
    style_header(ws, start_row)
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)


def build_excel_report() -> None:
    dso_segment = run_query(SQL_DIR / "dso_by_segment.sql")
    late_month = run_query(SQL_DIR / "late_payment_rate_by_month.sql")
    dso_overall = run_query(SQL_DIR / "dso_overall.sql")
    aging = run_query(SQL_DIR / "aging_buckets.sql")

    wb = Workbook()
    summary = wb.active
    summary.title = "Executive Summary"

    summary["A1"] = "Invoice & Cash Flow Analytics — Summary Report"
    summary["A1"].font = Font(size=16, bold=True, color="1F4E79")
    summary["A2"] = f"As of {AS_OF_DATE}"
    summary["A4"] = "Overall DSO (days)"
    summary["B4"] = float(dso_overall.iloc[0]["dso_days"])
    summary["A5"] = "Total Outstanding AR"
    summary["B5"] = float(dso_overall.iloc[0]["total_outstanding_ar"])
    summary["A6"] = "Total Invoices"
    summary["B6"] = int(dso_overall.iloc[0]["invoice_count"])

    for row in range(4, 7):
        summary[f"A{row}"].font = Font(bold=True)

    write_sheet(wb, "DSO by Segment", dso_segment)
    write_sheet(wb, "Late Rate by Month", late_month)
    write_sheet(wb, "Aging Buckets", aging)

    output = EXPORT_DIR / "cash_flow_summary.xlsx"
    wb.save(output)
    print(f"Saved Excel report -> {output}")


def main() -> None:
    export_tableau_csvs()
    build_excel_report()


if __name__ == "__main__":
    main()
